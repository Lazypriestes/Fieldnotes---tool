"""Import a question tree from a file.

    tldraw (.tldr)      board file; arrows (start -> end) and frames give the hierarchy
    Word (.docx)        headings -> topics, paragraphs / list levels -> questions & follow-ups
    PDF                 text layer (pypdf); a scanned / image-only PDF goes through OCR
    image               on-device OCR (macOS Vision, tools/ocr.swift) — photos, screenshots,
                        and Miro boards exported as PNG/JPG
    .txt / .md / .csv   markdown headings, bullets and indentation

extract() turns the file into Items (text + optional outline level + optional position).
build_tree() maps an outline straight onto topics/questions/follow-ups when the source has a
clear structure; otherwise the local LLM groups the lines. The LLM only answers with LINE
NUMBERS, so every question in the tree is the source's own wording, never a rewrite.

Result: [{"title": str, "questions": [{"text": str, "followups": [str, ...]}]}]
"""

import base64
import csv
import io
import json
import os
import re
import subprocess
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
OCR_SRC = os.path.join(HERE, "tools", "ocr.swift")
OCR_BIN = os.path.join(HERE, "tools", "ocr")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".heic", ".heif", ".webp", ".gif", ".tif", ".tiff", ".bmp"}
MAX_LINES = 300


class ImportError_(Exception):
    """A user-facing problem with the file (shown in the Import window)."""


class Item:
    __slots__ = ("text", "level", "x", "y", "h")

    def __init__(self, text, level=None, x=None, y=None, h=None):
        self.text, self.level, self.x, self.y, self.h = text.strip(), level, x, y, h


def _clean(s):
    return re.sub(r"\s+", " ", s or "").strip()


# --------------------------------------------------------------------------- readers
# bullets and numbering: '-', '•', '1.', '1)', '(a)', and dotted outlines with or without a
# trailing dot ('2.1', '2.1.') — a bare number like '2024' is NOT a bullet
_BULLET = re.compile(r"^(\s*)(?:[-*+•◦▪▫‣–]|\(?\d+(?:\.\d+)+[.)]?|\(?\d+[.)]|\(?[a-zA-Z][.)])\s+(.*)$")
_NUMDEPTH = re.compile(r"^\s*\(?(\d+(?:\.\d+)*)(?:[.)]|\s)")
_MD_HEAD = re.compile(r"^\s*(#{1,6})\s+(.*)$")


def from_text(s):
    """Markdown / plain text / PDF text. Levels come from headings, numbering depth
    ('2.1.' is deeper than '2.') and indentation (ranked, so any indent width works)."""
    lines = [l.rstrip() for l in s.replace("\r", "").split("\n")]
    indents = sorted({len(l.expandtabs(4)) - len(l.expandtabs(4).lstrip()) for l in lines if l.strip()})
    def indent_rank(l):
        n = len(l.expandtabs(4)) - len(l.expandtabs(4).lstrip())
        return max(i for i, v in enumerate(indents) if v <= n + 1)   # within 1 space counts as the same level
    items, head = [], -1
    structured = False
    for l in lines:
        if not l.strip():
            continue
        m = _MD_HEAD.match(l)
        if m:
            head = len(m.group(1)) - 1
            items.append(Item(m.group(2), head)); structured = True
            continue
        b = _BULLET.match(l)
        text = b.group(2) if b else l
        num = _NUMDEPTH.match(l) if b else None
        depth = (num.group(1).count(".") if num else 0) + indent_rank(l)
        if b or depth:
            structured = structured or depth > 0
        items.append(Item(text, head + 1 + depth))
    if not structured:                      # flat lines: let the LLM decide the grouping
        for it in items:
            it.level = None
    return items, structured


def from_docx(raw):
    from docx import Document
    doc = Document(io.BytesIO(raw))
    items, head, structured, title = [], -1, False, None
    for p in doc.paragraphs:
        text = _clean(p.text)
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style == "Title":
            title = title or text
            continue
        m = re.match(r"Heading (\d)", style)
        if m:
            head = int(m.group(1)) - 1
            items.append(Item(text, head)); structured = True
            continue
        ilvl = None
        numPr = p._p.pPr.numPr if p._p.pPr is not None else None
        if numPr is not None and numPr.ilvl is not None:
            ilvl = int(numPr.ilvl.val)
        elif style.startswith("List"):
            m2 = re.search(r"(\d)$", style)
            ilvl = int(m2.group(1)) - 1 if m2 else 0
        if ilvl is not None:
            structured = True
        items.append(Item(text, head + 1 + (ilvl or 0)))
    for t in doc.tables:                    # questions kept in a table: one item per cell
        for row in t.rows:
            for cell in row.cells:
                if _clean(cell.text):
                    items.append(Item(_clean(cell.text), head + 1))
    if not structured:
        for it in items:
            it.level = None
    title = title or (doc.core_properties.title or None)
    return items, structured, title


def from_pdf(raw):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    text = []
    for page in reader.pages[:40]:
        try:
            text.append(page.extract_text(extraction_mode="layout") or "")
        except Exception:
            text.append(page.extract_text() or "")
    joined = "\n".join(text)
    if len(re.sub(r"\s", "", joined)) < 20:          # scanned / image-only PDF -> OCR page 1
        with tempfile.TemporaryDirectory() as d:
            src, png = os.path.join(d, "in.pdf"), os.path.join(d, "page.png")
            open(src, "wb").write(raw)
            r = subprocess.run(["/usr/bin/sips", "-s", "format", "png", src, "--out", png],
                               capture_output=True, text=True)
            if r.returncode != 0 or not os.path.exists(png):
                raise ImportError_("this PDF has no text layer and could not be converted for OCR")
            return from_image(open(png, "rb").read()), False, "OCR of PDF page 1"
    items, structured = from_text(joined)
    return items, structured, None


def _ensure_ocr():
    if os.path.exists(OCR_BIN) and os.path.getmtime(OCR_BIN) >= os.path.getmtime(OCR_SRC):
        return
    r = subprocess.run(["swiftc", "-O", OCR_SRC, "-o", OCR_BIN], capture_output=True, text=True)
    if r.returncode != 0:
        raise ImportError_("could not build the on-device OCR helper (needs Xcode command-line tools: "
                           "xcode-select --install)")


def from_image(raw, ext=".png"):
    """OCR with positions. Text blocks that sit directly under each other with the same left
    edge are merged (a sticky note's lines), so each Item is one note / one question."""
    _ensure_ocr()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
        f.write(raw); path = f.name
    try:
        r = subprocess.run([OCR_BIN, path], capture_output=True, text=True, timeout=120)
    finally:
        os.unlink(path)
    if r.returncode != 0:
        raise ImportError_("could not read text from this image")
    blocks = []
    for line in r.stdout.splitlines():
        parts = line.split("\t", 4)
        if len(parts) == 5 and parts[4].strip():
            x, y, w, h = map(float, parts[:4])
            blocks.append([x, y, w, h, parts[4].strip()])
    blocks.sort(key=lambda b: (round(b[1], 2), b[0]))
    merged = []
    for b in blocks:
        for m in reversed(merged[-6:]):
            same_col = abs(m[0] - b[0]) < 0.015
            gap = b[1] - (m[1] + m[3])
            continues = not re.search(r"[?.!:]$", m[4]) or b[4][:1].islower()
            if same_col and -0.002 < gap < 1.1 * b[3] and continues:
                m[4] = m[4] + " " + b[4]
                m[3] = (b[1] + b[3]) - m[1]
                break
        else:
            merged.append(b)
    # h = height of ONE text line in the block (headings are set bigger than sticky text)
    return [Item(b[4], None, b[0], b[1], _line_h(b, blocks)) for b in merged]


def _line_h(block, raw_blocks):
    """smallest single-line height among the raw OCR lines that make up a merged block"""
    hs = [r[3] for r in raw_blocks if abs(r[0] - block[0]) < 0.015 and block[1] - 0.002 <= r[1] <= block[1] + block[3]]
    return min(hs) if hs else block[3]


def from_layout(items):
    """Boards and photos: headings are the BIGGER text (or ALL CAPS) without a question mark;
    each question joins the heading above it whose column it is in (nearest horizontally).
    Returns topics, or None when the layout has no recognisable headings."""
    qs = [it for it in items if it.x is not None and it.h]
    if len(qs) < 2:
        return None
    hs = sorted(it.h for it in qs if is_question(it.text)) or sorted(it.h for it in qs)
    body = hs[len(hs) // 2]
    heads = [it for it in qs if not is_question(it.text) and len(it.text) <= 60
             and (it.h >= 1.2 * body or (it.text.isupper() and len(it.text) > 2))]
    if not heads:
        return None
    groups = {id(h): [] for h in heads}
    stray = []
    for it in qs:
        if it in heads:
            continue
        above = [h for h in heads if h.y <= it.y + 0.005]
        if not above:
            stray.append(it); continue
        best = min(above, key=lambda h: (abs(h.x - it.x), it.y - h.y))
        groups[id(best)].append(it)
    topics = []
    for h in sorted(heads, key=lambda h: (round(h.y, 2), h.x)):
        members = [m for m in groups[id(h)] if is_question(m.text) or len(m.text.split()) >= 3]
        if members:
            members.sort(key=lambda m: (m.y, m.x))
            topics.append({"title": h.text, "questions": [{"text": m.text, "followups": []} for m in members]})
    loose = [s for s in stray if is_question(s.text)]
    if loose:
        topics.append({"title": "Other questions", "questions": [{"text": s.text, "followups": []} for s in loose]})
    return topics if sum(len(t["questions"]) for t in topics) >= 2 else None


def _tiptap_text(node):
    if isinstance(node, dict):
        if node.get("type") == "text":
            return node.get("text", "")
        parts = [_tiptap_text(c) for c in node.get("content", []) or []]
        sep = " " if node.get("type") in ("doc", "bulletList", "orderedList", "listItem") else ""
        return sep.join(p for p in parts if p)
    return ""


def from_tldraw(raw):
    """tldraw v2 (.tldr: {records: [...]}) — and the older v1 document format, best effort."""
    try:
        d = json.loads(raw.decode("utf-8"))
    except Exception:
        raise ImportError_("this .tldr file is not valid JSON")
    shapes, edges, frames = {}, [], set()
    recs = d.get("records")
    if recs is None and "document" in d:                         # tldraw v1
        for page in (d["document"].get("pages") or {}).values():
            for s in (page.get("shapes") or {}).values():
                if s.get("type") == "arrow":
                    continue
                t = _clean(s.get("text") or s.get("label") or "")
                if t:
                    pt = s.get("point") or [0, 0]
                    shapes[s["id"]] = {"text": t, "x": pt[0], "y": pt[1], "parent": None}
            ends = {}
            for b in (page.get("bindings") or {}).values():
                ends.setdefault(b.get("fromId"), {})[b.get("handleId")] = b.get("toId")
            for e in ends.values():
                if e.get("start") and e.get("end"):
                    edges.append((e["start"], e["end"]))
        recs = []
    by_id = {r.get("id"): r for r in recs or []}
    arrow_ends = {}
    for r in recs or []:
        if r.get("typeName") == "shape":
            p = r.get("props") or {}
            if r.get("type") == "arrow":
                for term in ("start", "end"):                      # pre-bindings arrows (tldraw <= 2.3)
                    v = p.get(term) or {}
                    if isinstance(v, dict) and v.get("type") == "binding" and v.get("boundShapeId"):
                        arrow_ends.setdefault(r["id"], {})[term] = v["boundShapeId"]
                continue
            if r.get("type") == "frame":
                frames.add(r["id"])
                text = _clean(p.get("name", ""))
            else:
                text = _clean(p.get("text") or _tiptap_text(p.get("richText")) or "")
            if text:
                shapes[r["id"]] = {"text": text, "x": r.get("x", 0), "y": r.get("y", 0), "parent": r.get("parentId")}
        elif r.get("typeName") == "binding" and r.get("type") == "arrow":
            term = (r.get("props") or {}).get("terminal")
            if term in ("start", "end"):
                arrow_ends.setdefault(r.get("fromId"), {})[term] = r.get("toId")
    for e in arrow_ends.values():
        if e.get("start") in shapes and e.get("end") in shapes and e["start"] != e["end"]:
            edges.append((e["start"], e["end"]))
    if not shapes:
        raise ImportError_("no text found on this tldraw board")

    parent = {}
    for a, b in edges:                                            # arrow start -> end = parent -> child
        parent.setdefault(b, a)
    for sid, s in shapes.items():                                 # inside a frame = child of the frame
        if sid not in parent and s["parent"] in shapes and s["parent"] in frames:
            parent[sid] = s["parent"]
    structured = bool(parent)
    xs = [s["x"] for s in shapes.values()]; ys = [s["y"] for s in shapes.values()]
    span = lambda v, lo, hi: (v - lo) / ((hi - lo) or 1)
    if not structured:
        return [Item(s["text"], None, span(s["x"], min(xs), max(xs)), span(s["y"], min(ys), max(ys)))
                for s in sorted(shapes.values(), key=lambda s: (s["y"], s["x"]))], False

    kids = {}
    for c, p in parent.items():
        kids.setdefault(p, []).append(c)
    order = lambda ids: sorted(ids, key=lambda i: (shapes[i]["y"], shapes[i]["x"]))
    items, seen = [], set()

    def walk(i, depth):
        if i in seen:                                              # arrows can form cycles
            return
        seen.add(i)
        items.append(Item(shapes[i]["text"], depth))
        for c in order(kids.get(i, [])):
            walk(c, depth + 1)

    for root in order([i for i in shapes if i not in parent]):
        walk(root, 0)
    return items, True


def from_csv(raw):
    text = raw.decode("utf-8", "replace")
    rows = list(csv.reader(io.StringIO(text)))
    return [Item(_clean(c)) for row in rows for c in row if _clean(c)]


def extract(name, raw):
    """-> (items, structured, title_hint, source_label)"""
    ext = os.path.splitext((name or "").lower())[1]
    if ext == ".tldr" or (ext == ".json" and b'"records"' in raw[:4000]):
        items, st = from_tldraw(raw); return items, st, None, "tldraw board"
    if ext == ".docx":
        items, st, title = from_docx(raw); return items, st, title, "Word document"
    if ext == ".doc":
        raise ImportError_("old .doc files aren't supported — save it as .docx (Word › Save As)")
    if ext == ".pdf":
        items, st, note = from_pdf(raw); return items, st, None, note or "PDF"
    if ext in IMAGE_EXT:
        return from_image(raw, ext), False, None, "image (on-device OCR)"
    if ext == ".csv":
        return from_csv(raw), False, None, "CSV"
    if ext in (".rtb",):
        raise ImportError_("Miro backups (.rtb) can't be read — in Miro use Export › Save as PDF or image")
    items, st = from_text(raw.decode("utf-8", "replace"))
    return items, st, None, "text"


# --------------------------------------------------------------------------- tree building
def is_question(t):
    return t.rstrip().endswith("?")


def _forest(items):
    """Items with levels -> nested nodes {text, children}, via a level stack."""
    roots, stack = [], []
    for it in items:
        node = {"text": it.text, "children": []}
        while stack and stack[-1][0] >= it.level:
            stack.pop()
        (stack[-1][1]["children"] if stack else roots).append(node)
        stack.append((it.level, node))
    return roots


def _descendants(n):
    for c in n["children"]:
        yield c["text"]
        yield from _descendants(c)


def _question(n):
    return {"text": n["text"], "followups": list(_descendants(n))}


def from_outline(items):
    roots = _forest(items)
    title = None
    # a single non-question root wrapping a deeper outline is the document title, not a topic
    while (len(roots) == 1 and not is_question(roots[0]["text"]) and roots[0]["children"]
           and any(c["children"] for c in roots[0]["children"])):
        title = title or roots[0]["text"]
        roots = roots[0]["children"]
    topics, loose = [], []
    for r in roots:
        if r["children"] and not (is_question(r["text"]) and not any(c["children"] for c in r["children"])):
            topics.append({"title": r["text"], "questions": [_question(c) for c in r["children"]]})
        elif r["children"] or is_question(r["text"]):
            loose.append(_question(r))                 # a question (with its follow-ups) at top level
        else:
            topics.append({"title": r["text"], "questions": []})
    if loose:
        topics.append({"title": "Questions" if not topics else "Other questions", "questions": loose})
    topics = [t for t in topics if t["questions"]] or topics
    return topics, title


LLM_SYS = (
    "You turn the text of an interview-question document or board into a question tree. "
    "Input: numbered lines; lines from an image or whiteboard also carry their position "
    "(x, y from 0 to 1; things close together usually belong together; a heading usually sits "
    "above or at the centre of its group). Group the lines into TOPICS. Each topic has "
    "QUESTIONS; a question may have FOLLOW-UPS (sub-questions or probes that belong to it). "
    "Rules: refer to lines ONLY by their number - never rewrite text. A topic's title is a "
    "heading line's number; only if a group has no heading line, write a short title of at "
    "most 5 words instead. Skip lines that are neither questions nor headings (page numbers, "
    "dates, names, logos, decoration). Keep the original order. Every question line is used "
    "at most once. Output ONLY JSON: {\"topics\": [{\"title_line\": <number or null>, "
    "\"title\": \"<only when title_line is null>\", \"questions\": [{\"line\": <number>, "
    "\"followups\": [<number>, ...]}]}]}"
)


def from_llm(items, llm_json):
    items = items[:MAX_LINES]
    lines = []
    for i, it in enumerate(items, 1):
        pos = f"(x={it.x:.2f}, y={it.y:.2f}) " if it.x is not None else ""
        lines.append(f"[{i}] {pos}{it.text}")
    out = llm_json(LLM_SYS, "\n".join(lines))
    used, topics = set(), []

    def take(n):
        if isinstance(n, int) and 1 <= n <= len(items) and n not in used:
            used.add(n)
            return items[n - 1].text
        return None

    for t in (out.get("topics") or []):
        if not isinstance(t, dict):
            continue
        title = take(t.get("title_line")) if t.get("title_line") is not None else None
        title = title or _clean(str(t.get("title") or ""))[:60] or "Topic"
        qs = []
        for q in t.get("questions") or []:
            if not isinstance(q, dict):
                continue
            text = take(q.get("line"))
            if not text:
                continue
            fus = [f for f in (take(n) for n in (q.get("followups") or [])) if f]
            qs.append({"text": text, "followups": fus})
        if qs:
            topics.append({"title": title, "questions": qs})
    # questions the model skipped are kept rather than silently lost
    missed = [items[i - 1].text for i in range(1, len(items) + 1)
              if i not in used and is_question(items[i - 1].text)]
    if missed:
        topics.append({"title": "Unsorted questions", "questions": [{"text": m, "followups": []} for m in missed]})
    return topics


def build_tree(name, raw, llm_json):
    """-> {"name", "topics", "info"}"""
    items, structured, title, source = extract(name, raw)
    items = [it for it in items if it.text]
    if not items:
        raise ImportError_(f"no text found in this {source}")
    topics, how = None, ""
    if structured:
        topics, t2 = from_outline(items)
        title = title or t2
        how = "kept the document's own structure"
        if sum(len(t["questions"]) for t in topics) < 2:           # outline too thin -> let the LLM try
            topics = None
    if topics is None and any(it.h for it in items):
        topics = from_layout(items)
        how = "grouped by the board's layout (headings and columns)"
    if topics is None:
        topics = from_llm(items, llm_json)
        how = "organized by the local AI (wording kept exactly as written)"
    if not any(t["questions"] for t in topics):
        raise ImportError_("couldn't find any questions in this file")
    stem = os.path.splitext(os.path.basename(name or "imported"))[0]
    nq = sum(len(t["questions"]) for t in topics)
    nf = sum(len(q["followups"]) for t in topics for q in t["questions"])
    return {"name": _clean(title or stem)[:60] or "Imported tree", "topics": topics,
            "info": f"From {source}: {len(topics)} topics, {nq} questions, {nf} follow-ups — {how}."}


def decode_upload(data):
    """'data:...;base64,XXXX' or bare base64 -> bytes"""
    if not isinstance(data, str) or not data:
        raise ImportError_("no file received")
    if data.startswith("data:"):
        data = data.split(",", 1)[1]
    try:
        return base64.b64decode(data)
    except Exception:
        raise ImportError_("the file could not be decoded")
